from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


doc = Document()
section = doc.sections[0]
section.top_margin = Pt(54)
section.bottom_margin = Pt(54)
section.left_margin = Pt(72)
section.right_margin = Pt(72)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("MoveIt2/RViz 执行抖动问题记录")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(16)

items = [
    ("问题现象", [
        "RViz 中 Plan and Execute 已显示成功，终端有 SUCCEEDED。",
        "灰色当前模型在初始位置和目标位置之间来回跳动/抽搐。",
        "ros2 topic info /joint_states -v 显示 Publisher count: 2。",
    ]),
    ("根本原因", [
        "/joint_states 同时被 fake_trajectory_controller 和 joint_state_publisher 发布。",
        "fake_trajectory_controller 发布执行后的关节角；joint_state_publisher 仍发布初始关节角。",
        "robot_state_publisher/RViz 收到两套关节状态，模型就在两套姿态之间跳变。",
    ]),
    ("排查方法", [
        "检查发布者：ros2 topic info /joint_states -v。",
        "检查残留节点：ros2 node list | grep joint_state_publisher。",
        "检查进程：ps -eo pid,ppid,cmd | grep joint_state_publisher | grep -v grep。",
        "判断执行是否成功，看 fake controller 是否收到轨迹，以及 MoveIt 是否输出 SUCCEEDED。",
    ]),
    ("解决办法", [
        "fake 执行模式下只保留 fake_trajectory_controller 发布 /joint_states。",
        "停止旧 launch 后清理残留：pkill -f joint_state_publisher。",
        "重新 source install/setup.bash 后启动 demo.launch.py。",
        "修改 launch 条件：fake_execution:=true 时禁止 joint_state_publisher 启动。",
    ]),
    ("经验总结", [
        "MoveIt 执行成功不等于 RViz 显示一定正确，显示异常优先检查 /joint_states。",
        "一个机器人模型通常只能有一个权威 JointState 发布源。",
        "遇到模型跳动、闪烁、抽搐，先查 topic publisher 数量，再查是否有旧节点残留。",
        "终端中的 unknown goal/result response 可先降级处理，核心判断看控制器完成状态和 SUCCEEDED。",
    ]),
]

for heading, bullets in items:
    p = doc.add_paragraph()
    r = p.add_run(heading)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(12)
    for bullet in bullets:
        bp = doc.add_paragraph(style=None)
        bp.paragraph_format.left_indent = Pt(18)
        bp.paragraph_format.first_line_indent = Pt(-10)
        br = bp.add_run("- " + bullet)
        br.font.name = "Arial"
        br.font.size = Pt(10.5)

doc.save("MoveIt2_RViz_joint_states_issue_note.docx")
